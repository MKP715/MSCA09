#!/usr/bin/env python3
"""Anonymity scan of the documents the website links to (Tradition Eleven). COUNTS ONLY.

Reads every row of the document index (data/documents/*.csv, or the older single file
data/documents.csv), opens the file from your Google Drive for Desktop copy of the Area Drive
(MSCA09AA/<drive_path>), extracts the text of EVERY page with two different PDF readers
(pypdf and PyMuPDF — each one finds text the other misses) and looks for:

  email     personal e-mail addresses (gmail/yahoo/…, or any address whose name part is not a
            service role). Role mailboxes (…@msca09aa.org, district/committee role addresses and
            everything in data/contact-allowlist.csv) are fine.
  phone     phone numbers that are not public office lines (data/central-offices.csv,
            data/contact-allowlist.csv, data/resources.csv, toll-free, GSO, Zoom dial-in numbers).
  sobriety  "sobriety date", "years sober", "sober anniversary", "aniversario de sobriedad",
            "sobriety birthday" … next to a date or a number.
  address   a street address that does not look like a meeting place (church, hall, club, office…).
  name      "First Lastname" next to a service role (DCMC, GSR, Chair, Delegate, Trustee, Secretary…),
            under "Present:" / "Attendance" / "Presentes:", or above a signature.
  unverified  pages that have no text layer (scanned images): nobody has read them yet.
            With --ocr (Windows 10/11) those pages are read with the Windows OCR engine first.

The report never prints what was found — only how many, per document. Look at the file
yourself (see scripts/README-documents.md, "Anonymity checklist").

  python scripts/scan-documents.py                  # every row with publish = yes
  python scripts/scan-documents.py --ids a,b        # only these rows (id column)
  python scripts/scan-documents.py --file x.pdf     # one local file (before you upload it)
  python scripts/scan-documents.py --redacted       # only the '-redacted' copies
  python scripts/scan-documents.py --since 2026-10-05T20:00   # only files added/changed since
  python scripts/scan-documents.py --all-rows       # also rows with publish = review / no
  python scripts/scan-documents.py --ocr            # OCR pages without text (Windows only)
  python scripts/scan-documents.py --hold           # move every flagged publish=yes row to
                                                    #   data/documents/held.csv (publish = review)
  python scripts/scan-documents.py --redact IN.pdf OUT.pdf   # make a cleaned copy of a PDF:
        # names → "First L.", e-mail/phone/address/sobriety date → "[removed]"; then re-scans it.

Exit code 1 when a publish=yes row has a finding (so a check can stop a deploy).
Needs: pip install --user pypdf pymupdf python-docx   (openpyxl / python-pptx optional)
Drive folder: set the environment variable MSCA09_DRIVE to your Google Drive for Desktop copy of the
Area folder (e.g. MSCA09_DRIVE="G:/My Drive/MSCA09AA"), or pass --drive. Not needed for --file / --redact.
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import csv
import datetime as dt
import hashlib
import io
import json
import logging
import os
import re
import subprocess
import sys
import tempfile
import unicodedata

logging.disable(logging.CRITICAL)  # pypdf is chatty about broken PDFs

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(REPO, "data")
CACHE_DIR = os.path.join(REPO, ".cache")
DRIVE = (os.environ.get("MSCA09_DRIVE") or "").rstrip("/\\")  # your local copy of the Area Drive folder (no default)
SCAN_VERSION = 16  # bump when the rules change, so cached results are recomputed

# ------------------------------------------------------------------ index files
def index_files():
    """The document index: data/documents/*.csv (split files) or data/documents.csv (older layout)."""
    d = os.path.join(DATA, "documents")
    if os.path.isdir(d):
        return sorted(os.path.join(d, f) for f in os.listdir(d) if f.lower().endswith(".csv"))
    f = os.path.join(DATA, "documents.csv")
    return [f] if os.path.exists(f) else []


def read_rows(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        r = csv.reader(fh)
        header = next(r)
        out = []
        for i, cells in enumerate(r, start=2):
            if not cells or not "".join(cells).strip() or cells[0].lstrip().startswith("#"):
                continue
            row = dict(zip(header, cells + [""] * (len(header) - len(cells))))
            row["_file"], row["_line"] = path, i
            out.append(row)
        return header, out


# ------------------------------------------------------------------ allow lists
def _digits(s):
    d = re.sub(r"\D", "", s or "")
    return d[1:] if len(d) == 11 and d.startswith("1") else d


def load_allow():
    emails, phones = set(), set()
    for name in ("contact-allowlist.csv", "central-offices.csv", "resources.csv"):
        p = os.path.join(DATA, name)
        if not os.path.exists(p):
            continue
        with open(p, encoding="utf-8-sig", newline="") as fh:
            for cells in csv.reader(fh):
                for c in cells:
                    for m in re.finditer(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", c):
                        if name == "contact-allowlist.csv":
                            emails.add(m.group(0).lower())
                    for m in re.finditer(r"\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}", c):
                        phones.add(_digits(m.group(0)))
    return emails, phones


# Public lines that are not in the CSV files: GSO switchboard/departments, Grapevine, Zoom dial-ins.
PUBLIC_PHONES = set(
    """2128703400 2128703023 2128703120 2128703019 2128703017 2128703110 2128703116 2128703351 2128703430
    2128703107 2128703302 2128703311 2128703337 2128703082 2128703400
    6699006833 3462487799 2532158782 3017158592 3126266799 9292056099 6694449171 7193594580 2532050468
    6892781000 3052241968 3092053325 3602095623 3863475053 5074734847 5642172000 6469313860 6699009128
    6468769923 6465588656 4086380968 4086380986""".split()
)
TOLL_FREE = ("800", "888", "877", "866", "855", "844", "833")

# ------------------------------------------------------------------ e-mail rules
EMAIL = re.compile(r"[A-Za-z0-9][A-Za-z0-9._%+\-]*\s?@\s?[A-Za-z0-9\-]+(?:\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}")
FREE_MAIL = re.compile(
    r"(^|\.)(gmail|googlemail|yahoo|ymail|rocketmail|icloud|me|mac|hotmail|outlook|live|msn|aol|att|sbcglobal|"
    r"pacbell|bellsouth|comcast|cox|verizon|earthlink|charter|frontier|roadrunner|rr|twc|juno|netzero|protonmail|"
    r"proton|gmx|mail|zoho|q|prodigy|mindspring|adelphia|netscape|excite|lycos|peoplepc|socal|dslextreme)\.",
    re.I,
)
AA_DOMAINS = re.compile(
    r"(msca09aa\.org|(^|\.)aa\.org|aagrapevine\.org|lavina\.org|aalavina\.org|praasa\.org|pacificregion|area09\.org|"
    r"area9|area-?\d{2}|oc-aa\.org|aainlandempire|aasbco|asbco|hacoaa|aainthedesert|aa-intergroup|intergroup|"
    r"intergrupo|centraloffice|oficinacentral|scaa|ypaa|aaws|aahistory|alcoholics-anonymous|nagsc|district\d|"
    r"aadistrict|distrito)",
    re.I,
)
ROLE_WORDS = sorted(
    """area msca district distrito dist panel delegate delegado delegada alternate alterno alterna alt chairman
    chairperson chairwoman cochair chair coordinador coordinadora coordinator secretary secretario secretaria
    treasurer tesorero tesorera treas registrar registrador registradora archivist archives archivos archive
    literature literatura webmaster webservant website web tech technology tecnologia committee comite information
    info office oficina contact contacto admin help hotline events event registration newsletter editor corrections
    correccionales correcciones treatment tratamiento grapevine lavina vina accessibilities accessibility
    accesibilidad remote communities convention conventions ypaa intergroup intergrupo centraloffice central
    contributions contribuciones finance finanzas forum foro praasa assembly asamblea service services servicio
    servicios gso aaws translation traduccion outreach bridging btg sponsorship cpc pi hni handi hi dcmc dcm mcd
    cmcd gsr rsg sober aa of the and de la el los las mid southern sur medio california calif ca socal oc hispanic
    hispano hispana spanish english public publica publicinfo communications comunicaciones liaison enlace group
    grupo groups grupos chairs reps rep sec pres vice co hospitality hospitalidad schedule cec special needs
    fellowship recovery unity audit sevtrad seventh tradition corr gv lv mtg meeting meetings speakers speaker
    workshop taller school escuela trustee custodio region regional pacific orange county riverside sanbernardino
    longbeach southbay harbor desert inland empire noreply donotreply postmaster mailer""".split(),
    key=len,
    reverse=True,
)
_ROLE_TOK = re.compile("|".join(re.escape(t) for t in ROLE_WORDS if len(t) >= 4 or t in ("dcm", "gsr", "rsg", "mcd", "cpc", "cec", "gso", "btg", "hni", "sec", "alt", "rep", "web", "msca")))
_FILLER = re.compile(r"(?:aa|of|the|and|de|del|la|el|los|las|ca|oc|co|hi|pi|mid|sur|ie|sb|lb|sd|us|org|ap|ar|p|d|v|y|a|e|s)*")


def email_is_personal(addr, allow_emails):
    a = re.sub(r"\s", "", addr).lower().strip(".,;:<>()[]")
    if a in allow_emails:
        return False
    local, _, dom = a.partition("@")
    if AA_DOMAINS.search(dom) or dom.endswith((".gov", ".edu", ".mil", ".us")):
        return False
    if dom in ("example.com", "domain.com", "email.com", "yourdomain.com"):
        return False
    # Whatever is left of the name part after removing service words, digits and dots:
    # "district18dcmc" → "" (a role mailbox); "jsmith", "maryg.dcm" → "jsmith", "maryg" (a person).
    rest = re.sub(r"^p(?=\d)", "", local)
    rest = re.sub(r"[\d._+\-]", "", rest)
    rest = _ROLE_TOK.sub("", rest)
    if len(rest) < 3 or _FILLER.fullmatch(rest):
        return False
    return True  # free-mail or any other domain: a person's own address


def join_split_emails(t):
    """PDF text often breaks an address over two lines ("name@yahoo.c" + "om"): join the pieces."""
    t = re.sub(
        r"(@[A-Za-z0-9\-]+(?:\.[A-Za-z0-9\-]*)*)[ \t]*\n[ \t]*([A-Za-z0-9.\-]*[A-Za-z])(?=\W|$)",
        lambda m: m.group(1) + m.group(2) if "." in m.group(1) + m.group(2) and not re.search(r"\.[a-z]{2,}$", m.group(1)) else m.group(0),
        t,
    )
    t = re.sub(r"([A-Za-z0-9._%+\-]+)[ \t]*\n[ \t]*(@[A-Za-z0-9\-]+\.)", r"\1\2", t)
    return t


# ------------------------------------------------------------------ phone rules
# Separators: hyphen (also the Unicode hyphens and dashes PDFs use) or dot, with spaces or a line break around them,
# or just spaces / a line break: "(949)\n555-0110", "714-\n\n555-0123", "909‐555‐0147", "805-5550199", "R.(714)555-0162".
_PSEP = r"[.\-\u2010-\u2015\u2212]"
PHONE = re.compile(
    r"(?<![\d\-$#\u2010-\u2015])(?<!\d\.)(?:\+?1[\s.\-]?)?"
    r"(?:\(\s?([2-9]\d{2})\s?\)\s{0,3}|([2-9]\d{2})(?:\s{0,3}" + _PSEP + r"\s{0,3}|/|\s{1,3}))"
    r"([2-9]\d{2})(?:\s{0,3}" + _PSEP + r"\s{0,3}|\s{1,3})?\d{4}(?!\d)(?!" + _PSEP + r"\d)"
)
PHONE_SKIP_BEFORE = re.compile(
    # "Zoom: 457 182 8915", "Zoom Meeting ID 457…", "zoom.us/j/457…" — the word Zoom right before the number
    r"(?:zoom(?:\.us/\w+/|\s*(?:meeting|reuni[oó]n))?(?:\s*(?:id|i\.d\.|#|number|n[uú]mero))?\s*[:#=-]?\s*$)|"
    # other labels a little further away: "ID de la reunión ad hoc: 6516089134", "Account # 4017400154"
    r"(?:(meeting\s*id|\bid\b|\bid\s*[:#]|i\.d\.|passcode|password|pass\s*code|contrase|c[oó]digo|\bpin\b|\bext\b|"
    r"conference\s*id|webinar|account|acct|routing|check\s*(no|#)|invoice|receipt|\bein\b|tax\s*id|order\s*(no|#)|"
    r"\$|isbn|item|catalog|cat\.?\s*no)[^\n]{0,25}$)",  # not 40: "on Zoom. Call or text Pat at 714…" is a phone
    re.I,
)


def phone_hits(text, allow_phones):
    out = set()
    for m in PHONE.finditer(text):
        d = _digits(m.group(0))[-10:]
        if len(d) != 10:
            continue
        if d in allow_phones or d in PUBLIC_PHONES or d[:3] in TOLL_FREE or d.startswith("212870"):
            continue
        if PHONE_SKIP_BEFORE.search(text[max(0, m.start() - 50) : m.start()]):
            continue
        out.add((d, m.group(0).strip()))
    return out


# ------------------------------------------------------------------ sobriety rules
SOBRIETY = re.compile(
    r"(sobriety\s+dates?|sobriety\s+birthdays?|sobriety\s+anniversar(?:y|ies)|sober\s+anniversar(?:y|ies)|"
    r"sober\s+birthdays?|years?\s+(?:of\s+)?sober(?:iety)?\b|yrs\.?\s+sober|sober\s+since|clean\s+and\s+sober\s+since|"
    r"aniversario\s+de\s+sobriedad|aniversarios\s+de\s+sobriedad|fecha\s+de\s+sobriedad|cumplea[nñ]os\s+de\s+sobriedad|"
    r"a[nñ]os\s+de\s+sobriedad|a[nñ]os\s+sobri[oa]s?|aa\s+birthday|anniversary\s+of\s+(?:his|her|my)\s+sobriety|"
    r"dates?\s+of\s+(?:my\s+|his\s+|her\s+|their\s+)?sobriety|sober\s+dates?|clean\s+dates?|dry\s+dates?|"
    r"got\s+sober\s+(?:on|in)|sobriety\s+began|sobri[oa]s?\s+desde|fecha\s+de\s+(?:mi\s+|su\s+)?sobriedad|"
    r"(?:mi|su)\s+sobriedad\s+(?:es|comenz[oó]|empez[oó])|aniversario\s+de\s+(?:mi\s+|su\s+)?sobriedad|"
    r"cumplo\s+\w+\s+a[nñ]os\s+(?:de\s+)?sobri)",
    re.I,
)
MONTH_RX = (
    r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sept?(?:ember)?|oct(?:ober)?|"
    r"nov(?:ember)?|dec(?:ember)?|enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|setiembre|octubre|"
    r"noviembre|diciembre)"
)
# Longest forms first, so "January 20, 2013" is matched whole (not just "January 20").
DATEISH = re.compile(
    r"(\b" + MONTH_RX + r"\.?\s+\d{1,2}(?:st|nd|rd|th)?,?\s+(?:de\s+|del\s+)?(?:19|20)?\d{2}\b"
    r"|\b(?:el\s+)?\d{1,2}\s+(?:de\s+)?" + MONTH_RX + r",?\s+(?:de\s+|del\s+)?(?:19|20)\d{2}\b"
    r"|\b\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4}\b"
    r"|\b" + MONTH_RX + r"\.?\s+\d{1,2}\b|\b\d{1,2}\s+(?:de\s+)?" + MONTH_RX
    + r"\b|\b" + MONTH_RX + r"\.?,?\s+(?:de\s+|del\s+)?(?:19|20)\d{2}\b|\b(?:19|20)\d{2}\b)",
    re.I,
)
NUMBERISH = re.compile(r"\b\d{1,2}\b|\b(?:one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|fifteen|twenty|thirty|forty|fifty|un|dos|tres|cuatro|cinco|seis|siete|ocho|nueve|diez|veinte|treinta|cuarenta)\b", re.I)


DATE_PHRASE = re.compile(r"date|fecha|birthday|cumplea|anniversar|aniversario", re.I)
REQUIREMENT = re.compile(
    r"requir|requisit|m[ií]nim|at least|por lo menos|al menos|suggest|suger|should have|debe(?:r[aá])?\s+tener|"
    r"continu|qualif|eligib|calific|recomend|recommend|desirable|deseable|preferably|preferible",
    re.I,
)
PERSONISH = re.compile(r"\b[A-Z][a-záéíóúñ]+\s[A-Z]\.(?=[\s,;)]|$)|\bmy name is\b|\bmi nombre es\b|\bme llamo\b|\bsoy alcoh|\bI'?m an alcoholic\b", re.I)
PRONOUN_BEFORE = re.compile(r"\b(I|I've|I'm|I have|I am|me|he|she|he's|she's|has been|have been|been|estoy|llevo|tengo|tiene)\b[^.\n]{0,25}$", re.I)


def sobriety_hits(text):
    """Spans (start, end) of a member's sobriety date, or of a member's years of sobriety next to a name.

    Flagged:  "My sobriety date is May 3, 1998", "fecha de sobriedad: 3/5/98", "Pat Q. (12 years sober)",
              "I've been sober since 2009".
    Not flagged: requirements ("five years of continuous sobriety"), anonymous sharing without a
              name or date ("tengo 30 años de sobriedad"), history ("living sober since 1935").
    """
    out = []
    for m in SOBRIETY.finditer(text):
        before = text[max(0, m.start() - 40) : m.start()]
        after = text[m.end() : m.end() + 45]
        phrase = m.group(0).lower()
        near = text[max(0, m.start() - 110) : m.end() + 90]
        span = None
        if DATE_PHRASE.search(phrase):
            dm = DATEISH.search(after)
            if dm and not re.search(r"[.;]\s+[A-Z]", after[: dm.start()]):
                span = (m.start(), m.end() + dm.end())
            else:
                db = list(DATEISH.finditer(before[-30:]))
                if db and not re.search(r"\d", phrase):
                    span = (m.start() - len(before[-30:]) + db[-1].start(), m.end())
        elif "since" in phrase:
            dm = DATEISH.match(after.lstrip()) or DATEISH.search(after[:16])
            if dm and PRONOUN_BEFORE.search(before) and not re.search(r"living\s*$", before, re.I):
                span = (m.start(), m.end() + after.find(dm.group(0)) + len(dm.group(0)))
        elif re.search(r"desde|got sober|began|comenz|empez|sobriedad\s+es", phrase):
            dm = DATEISH.search(after[:28])
            if dm:
                span = (m.start(), m.end() + dm.end())
        else:  # "N years sober", "N años de sobriedad"
            nm = list(NUMBERISH.finditer(before[-18:]))
            if nm and not REQUIREMENT.search(text[max(0, m.start() - 90) : m.end() + 60]) and (PERSONISH.search(near) or ROLE_NEAR.search(near)):
                span = (m.start() - len(before[-18:]) + nm[-1].start(), m.end())
        if span:
            out.append(span)
    return out


# ------------------------------------------------------------------ street addresses
STREET_ABBR = r"St|Ave|Av|Blvd|Rd|Dr|Ln|Ct|Cir|Pl|Pkwy|Hwy|Ter|Trl|Plz|Sq"
STREET_WORDS = (
    r"Street|Avenue|Boulevard|Road|Drive|Lane|Way|Court|Circle|Place|Parkway|Highway|Terrace|Trail|Loop|Square|"
    r"Calle|Avenida|Camino|Paseo"
)
_STREET_NAME = r"(?:\d{1,3}(?:st|nd|rd|th)|[A-Z][A-Za-z'\-]+|[A-Z]{2,})"
# "123 Main Ave", "456 W 1st Street", "789 Example Cir", "10 SAMPLE ST." — street name capitalised; the
# abbreviation (St, Ave, Blvd …) in any case, the full word (Street, Avenue …) capitalised.
ADDRESS = re.compile(
    r"(?<![\d$.,/#\-])\b(\d{1,5})\s+(?:[NSEW]\.?\s+|North\s+|South\s+|East\s+|West\s+)?((?:" + _STREET_NAME + r"\.?\s+){1,3})"
    r"(?:(?i:" + STREET_ABBR + r")\b\.?|(?:" + STREET_WORDS + r"|" + STREET_WORDS.upper() + r")\b)"
    r"(?:,?\s*(?:#|Apt\.?|Unit|Spc\.?|Space|Ste\.?|Suite)\s*[\w-]+)?"
)
# "12 sample ave, Apt C\nAnytown, CA" — a numbered line followed by a California city (any case street).
ADDRESS_CITY = re.compile(
    r"(?<![\d$.,/#\-])\b(\d{1,5})\s+([A-Za-z0-9][A-Za-z0-9'.\- ]{1,38}?)\s*(?:,|\n)\s*(?:(?:#|Apt\.?|Unit|Ste\.?|Suite)\s*[\w-]+\s*,?\s*)?"
    r"[A-Z][A-Za-z .]{2,28},\s*(?:CA|Ca\.|Calif\.?|California)\b\.?(?:,?\s*9\d{4})?"
)
# Words that show an address is a meeting place, office or public building (checked close to the address),
# including meeting lists ("Mon 7:00 PM Open Discussion").
VENUE_WORDS = re.compile(
    r"church|iglesia|chapel|capilla|parish|parroquia|temple|templo|cathedral|mission|misi[oó]n|lutheran|methodist|"
    r"presbyterian|episcopal|catholic|cat[oó]lica|baptist|unitarian|congregational|synagogue|our lady|nuestra se|"
    r"\bst\.? \w+'?s\b|saint|sacred|holy|\bhall\b|sal[oó]n|center|centre|centro|\bclub\b|alano|fellowship hall|"
    r"central office|oficina central|intergroup|intergrupo|hotel|\binn\b|resort|marriott|hilton|hyatt|sheraton|"
    r"doubletree|embassy suites|college|university|library|biblioteca|\bpark\b|parque|auditorium|auditorio|"
    r"hospital|medical|clinic|recreation|senior|ymca|ywca|veterans|legion|vfw|elks|moose|lodge|restaurant|"
    r"restaurante|caf[eé]|grill|convention|fairgrounds|stadium|plaza|mall|venue|held at|se llevar[aá]|"
    r"localizad|located at|ubicad|meets at|se re[uú]ne|directions|parking|estacionamiento|facility|institution|"
    r"prison|jail|correctional|rehab|camp\b|retreat|archives|archivos|repositor|historic|hist[oó]ric|onto\b|"
    r"turn (?:left|right)|\bexit\b|miles|millas|gire|salida|freeway|fwy|general service office|\bgso\b|\bosg\b|"
    r"\b\d{1,2}(?::\d{2})?\s*(?:a\.?m\.?|p\.?m\.?)(?![a-z])|\bnoon\b|mediod[ií]a|\b(?:mon|tue|tues|wed|thu|thur|thurs|fri|sat|sun)(?:day)?\b|"
    r"\b(?:lunes|martes|mi[eé]rcoles|jueves|viernes|s[aá]bado|domingo)\b|\bopen\b|\bclosed\b|abierta|cerrada|speaker|"
    r"step study|big book|discussion|men'?s|women'?s|stag|candlelight|meeting place|lugar de reuni[oó]n",
    re.I,
)
MAILING_LABEL = re.compile(
    r"(?:(?:area\s+)?dist(?:rict|rito)?\.?\s*\d{1,2}(?:\s*(?:&|and|y)\s*\d{1,2})?[ \t]*"
    r"|(?:treasurer|tesorer[oa]|mail\s+(?:to|checks?)|env[ií]e\s+a|send\s+(?:checks?|contributions?)\s+to|make\s+checks?\s+payable)[^\n]{0,25})"
    r"(?:[ \t]*\n){1,3}[^\n]{0,30}$",
    re.I,
)
VENUE_IN_NAME = re.compile(r"\b(?:school|escuela|church|iglesia|college|hospital|center|centre|centro|hall|club|library|park|plaza|hotel|inn)\b", re.I)
PERSONAL_ADDRESS_WORDS = re.compile(
    r"home of|my home|my house|residence|residencia|en casa de|casa de\b|home address|my address|mi direcci[oó]n|"
    r"send (it )?to me|mail to me|\bindividual\b|\bapt\.?\s*\w|apartment|\bspc\b|\bspace\s+\d",
    re.I,
)
# Public addresses that are not meeting places: the General Service Office, A.A. historic sites.
PUBLIC_ADDRESSES = {"475 riverside", "182 clinton", "855 ardmore"}
ARTICLE_STREET = re.compile(r"^(?:el|la|los|las|the|de|del|a|an|and|of|in|on|at|y|e|for|to|by|with|from|our|my|your|their)\.?\s*$", re.I)
NOT_STREET = re.compile(
    r"\b(?:years?|a[nñ]os|members?|miembros|people|personas|groups?|grupos|minutes|minutos|days?|d[ií]as|hours?|horas|"
    r"times?|veces|percent|copies|books?|libros|box|apartado|am|pm|a\.m|p\.m|noon|via|zoom|online|virtual|hybrid|"
    r"meeting|meetings|reuni[oó]n|steps?|traditions?|concepts?|panel|district|distrito|area|asa|asc|csa|gsr|rsg|dcm|mcd)\b",
    re.I,
)


def load_venues():
    """Street numbers + first street word of known meeting places (calendar, districts, central offices)."""
    keys = set()
    srcs = [os.path.join(DATA, f) for f in ("calendar.ics", "districts.csv", "central-offices.csv", "committees.csv", "resources.csv")]
    for p in srcs:
        if not os.path.exists(p):
            continue
        try:
            t = open(p, encoding="utf-8-sig", errors="ignore").read()
        except OSError:
            continue
        t = t.replace("\\,", ",").replace("\\n", "\n")
        for rx in (ADDRESS, ADDRESS_CITY):
            for m in rx.finditer(t):
                keys.add(_addr_key(m))
    return keys


def _addr_key(m):
    return m.group(1) + " " + (m.group(2).split() or [""])[0].lower().strip(".,")


def addr_hash(key):
    """A short one-way hash of "123 main": lets a whole run count documents per address without storing it."""
    return hashlib.sha1(key.encode("utf-8")).hexdigest()[:12]


def address_hits(text, venues, common=frozenset()):
    """(start, end, key, personal) of street addresses that do not look like meeting places.
    `common`: hashes of addresses found in many documents (meeting places, offices) — skipped unless the text
    around says it is someone's home."""
    out = []
    seen = []
    for rx in (ADDRESS, ADDRESS_CITY):
        for m in rx.finditer(text):
            if any(m.start() < b and a < m.end() for a, b in seen):
                continue
            words = m.group(2).strip()
            key = _addr_key(m)
            if key in venues or key in PUBLIC_ADDRESSES or ARTICLE_STREET.match(words) or NOT_STREET.search(words):
                continue
            if VENUE_IN_NAME.search(words):
                continue  # "Walton Intermediate School, Garden Grove, CA"
            if rx is ADDRESS_CITY and (not re.search(r"[A-Za-z]{3,}", words) or re.search(ROLE_TITLE, words)):
                continue
            if re.search(r"(box|apartado|p\.?\s?o\.?)\s*#?\s*$", text[max(0, m.start() - 14) : m.start()], re.I):
                continue  # P.O. Box 51446 — post-office boxes are fine
            if re.search(r"(?:\bCA|Calif\.?|California)[ ,.]*$", text[max(0, m.start() - 14) : m.start()]):
                continue  # the number is the ZIP code of the line before
            seen.append((m.start(), m.end()))
            near = text[max(0, m.start() - 90) : m.end() + 60]
            personal = bool(PERSONAL_ADDRESS_WORDS.search(near))
            # A street address used as a district's or committee's mailing address ("Mid So Cal Area Dist 5" on the
            # line above) is usually a member's home: never treated as a meeting place.
            if MAILING_LABEL.search(text[max(0, m.start() - 60) : m.start()]):
                personal = True
            if not personal and (VENUE_WORDS.search(near) or addr_hash(key) in common):
                continue
            out.append((m.start(), m.end(), key, personal))
    return out


# ------------------------------------------------------------------ names
FIRST_NAMES = set(
    """Aaron Abel Abigail Abraham Ada Adam Adan Adela Adolfo Adrian Adriana Agustin Agustina Aida Aileen Alan Alana
    Albert Alberto Alda Alec Alejandra Alejandro Alex Alexa Alexander Alexandra Alexis Alfonso Alfred Alfredo Alice
    Alicia Alina Alison Allan Allen Allison Alma Alonso Alvaro Alyssa Amanda Amber Amelia Amparo Amy Ana Anabel
    Andre Andrea Andres Andrew Andy Angel Angela Angelica Angelina Angie Anita Ann Anna Anne Annette Annie Anthony
    Antonia Antonio April Araceli Ariel Arlene Armando Arnold Arnulfo Arthur Arturo Ashley Audrey Aurelio Aurora
    Austin Barbara Barry Beatriz Becky Belinda Ben Benjamin Bernard Bernardo Bernice Bert Beth Betsy Betty Beverly
    Bianca Bill Billy Blanca Bob Bobbi Bobby Bonnie Brad Bradley Brandon Brenda Brent Brett Brian Bridget Brittany
    Brooke Bruce Bryan Bud Byron Caleb Camila Candace Candy Carl Carla Carlos Carmen Carol Carole Carolina Caroline
    Carolyn Carrie Casey Catalina Catherine Cathy Cecilia Cesar Chad Charlene Charles Charlie Charlotte Chelsea
    Cheri Cheryl Chris Christian Christina Christine Christopher Chuck Cindy Claire Clara Claudia Clayton Cliff
    Clifford Clint Clyde Cody Colleen Connie Conrad Consuelo Corey Courtney Craig Cris Cristian Cristina Crystal
    Curtis Cyndi Cynthia Dale Damian Dan Dana Daniel Daniela Danielle Danny Darla Darlene Darrell Darren Daryl Dave
    David Dawn Dean Deanna Debbie Deborah Debra Delia Denise Dennis Derek Diana Diane Dianne Diego Dolores Dominic
    Don Donald Donna Dora Doris Dorothy Doug Douglas Duane Dustin Dwight Earl Ed Eddie Edgar Edith Edna Eduardo
    Edward Edwin Efrain Elaine Eleanor Elena Eli Elias Elisa Elizabeth Ella Ellen Elmer Eloy Elsa Elvia Elvira
    Emilio Emily Emma Enrique Eric Erica Erik Erika Erin Ernest Ernesto Esperanza Esteban Estela Esther Ethan Eugene
    Eva Evelyn Fabian Faith Federico Felipe Felix Fernando Fidel Flor Florence Frances Francine Francisco Frank
    Frankie Fred Freddie Freddy Frederick Gabriel Gabriela Gail Gary Gene Genaro George Georgia Gerald Geraldine
    Gerardo Gilbert Gilberto Gina Ginger Glen Glenda Glenn Gloria Gordon Grace Graciela Greg Gregg Gregorio Gregory
    Guadalupe Guillermo Gustavo Hal Hank Harold Harry Harvey Hazel Heather Hector Heidi Helen Henry Herbert Heriberto
    Herman Hilda Hilario Holly Howard Hugo Humberto Ian Ignacio Ilene Inez Irene Iris Irma Isaac Isabel Isidro Israel
    Ivan Jack Jackie Jacob Jacqueline Jaime Jake James Jamie Jan Jane Janet Janice Janie Jared Jason Javier Jay
    Jean Jeanette Jeanne Jeff Jeffery Jeffrey Jenna Jennifer Jenny Jeremy Jerome Jerry Jesse Jessica Jessie Jesus
    Jill Jim Jimmy Jo Joan Joann Joanne Joaquin Jodi Jody Joe Joel Joey John Johnny Jon Jonathan Jorge Jose Josefina
    Joseph Josh Joshua Joy Joyce Juan Juana Juanita Judith Judy Julia Julian Julie Julio Justin Karen Kari Karl Karla
    Kate Katherine Kathleen Kathryn Kathy Katie Kay Keith Kelly Ken Kenneth Kenny Kerry Kevin Kim Kimberly Kirk Kris
    Kristen Kristin Kristina Kurt Kyle Lance Larry Laura Lauren Laurie Lawrence Leah Lee Leon Leonard Leonardo
    Leroy Leslie Leticia Lidia Lillian Lily Linda Lindsay Lisa Liz Lois Lola Lonnie Lora Lorena Lorenzo Loretta Lori
    Lorraine Louis Louise Lourdes Lucia Lucy Luis Luisa Lupe Luz Lydia Lynda Lynn Lynne Mabel Madeline Maggie
    Manuel Marc Marcela Marco Marcos Marcus Margaret Margarita Margie Maria Mariana Maribel Marie Marilyn Mario
    Marion Marisa Marisol Marissa Marjorie Mark Marlene Marsha Marta Martha Martin Marty Marvin Mary Maryann Mateo
    Matt Matthew Maureen Mauricio Max Maxine Megan Mel Melanie Melinda Melissa Melvin Mercedes Michael Micheal
    Michele Michelle Miguel Mike Mildred Miriam Missy Misty Mitch Mitchell Moises Monica Monique Morris Myra Myrna
    Nancy Naomi Natalie Natasha Nathan Neil Nelson Nestor Nicholas Nick Nicole Noe Noel Nora Norma Norman Octavio
    Olga Oliver Olivia Omar Oscar Pablo Pam Pamela Pat Patrice Patricia Patrick Patsy Patty Paul Paula Pauline
    Pedro Peggy Penny Perla Pete Peter Phil Philip Phillip Phyllis Pilar Priscilla Rachel Rafael Ralph Ramiro Ramon
    Ramona Randall Randy Raquel Raul Ray Raymond Rebecca Regina Reginald Rene Renee Reyna Rhonda Ricardo Rich
    Richard Rick Ricky Rigoberto Rita Rob Robbie Robert Roberta Roberto Robin Rocio Rodney Rodolfo Rodrigo Roger
    Rogelio Roland Rolando Ron Ronald Ronnie Rosa Rosalie Rosalinda Rosario Rose Rosemary Ross Roxanne Ruben Ruby
    Rudy Russ Russell Ruth Ryan Sabrina Sal Sally Salvador Sam Samantha Samuel Sandra Sandy Santiago Sara Sarah
    Scott Sean Sergio Seth Shane Shannon Sharon Shawn Sheila Shelley Sherri Sherry Shirley Sidney Silvia Simon
    Socorro Sofia Sonia Sonya Stacey Stacy Stan Stanley Stella Stephanie Stephen Steve Steven Stuart Sue Susan
    Susana Suzanne Sylvia Tamara Tammy Tanya Tara Ted Teresa Teri Terrance Terrence Terri Terry Thelma Theresa
    Thomas Tiffany Tim Timothy Tina Toby Todd Tom Tomas Tommy Toni Tony Tracey Traci Tracy Travis Trevor Trina Troy
    Tyler Valentin Valerie Vanessa Vera Verna Veronica Vicente Vicki Vickie Victor Victoria Vince Vincent Viola
    Violet Virginia Vivian Wade Walter Wanda Warren Wayne Wendy Wesley Willard William Willie Wilma Xavier Yesenia
    Yolanda Yvette Yvonne Zachary Zach""".split()
)
# Capitalised words that follow a first name but are not surnames (places, AA words, common words).
NOT_SURNAME = set(
    """Area District Distrito Committee Comite Comité Panel Zoom Service Servicio Servicios General Generales
    Conference Conferencia Delegate Delegado Delegada Alternate Suplente Chair Chairperson Secretary Secretario
    Secretaria Treasurer Tesorero Tesorera Registrar Registrador Grapevine Vina Viña Archives Archivos The Long Beach
    Santa Ana North South East West Central Office Oficina Church Iglesia Avenue Ave Street Blvd Boulevard Road
    Drive Way Park Valley Hills Hall Club Center Centro Room Sala January February March April May June July August
    September October November December Enero Febrero Marzo Abril Mayo Junio Julio Agosto Septiembre Octubre
    Noviembre Diciembre Monday Tuesday Wednesday Thursday Friday Saturday Sunday Lunes Martes Miercoles Miércoles
    Jueves Viernes Sabado Sábado Domingo County Condado Island Point Lake Mountain Springs City Ciudad Desert College
    University School Escuela Hispanic Hispano Spanish English Ingles Inglés Inter Intergroup Intergrupo Hotel
    Community Comunidad Fellowship Alano Foundation Unity Unidad Hope Serenity Sobriety Recovery Pacific Pacifico
    Pacífico Regional Region Forum Foro Trustee Custodio Board Junta Workshop Taller Assembly Asamblea Meeting
    Reunion Reunión Group Grupo Groups Grupos Report Informe Minutes Actas Motion Mocion Moción Budget Presupuesto
    Literature Literatura Public Information Treatment Corrections Cooperation Professional Technology Registration
    Finance Guidelines Policy Policies Accessibilities Remote Communities Convention Liaison Young People
    Correspondence Communications Bridging Gap Hospitals Institutions Elder Seventh Tradition Traditions Tradiciones
    Steps Pasos Concepts Conceptos Big Book Libro Grande Twelve Doce Living Sober Came Believe As Bill Sees It Daily
    Reflections Language Heart Members Member Miembro Miembros Newcomer Newcomers Sponsor Sponsorship Speaker
    Speakers Panel Agenda Item Items Background Action Actions Advisory Floor Discussion Report Reports Highlights
    Sharing Session Delegates Officers Officer Chairs Coordinator Coordinador Coordinadora Representative
    Representante Representatives Mid Southern California Riverside Orange Bernardino Diego Clemente Juan Jacinto
    Marcos Pedro Monica Barbara Clara Maria Luis Angeles Cruz Rosa Ana Fe Island Catalina Irvine Anaheim Fullerton
    Huntington Newport Costa Mesa Laguna Niguel Hills Woods Viejo Mission Torrance Carson Gardena Lakewood Downey
    Norwalk Whittier Cerritos Bellflower Compton Paramount Hemet Temecula Murrieta Menifee Perris Corona Norco
    Ontario Fontana Rialto Redlands Yucaipa Banning Beaumont Palm Desert Indio Coachella Victorville Hesperia Apple
    Barstow Needles Blythe Tustin Garden Grove Westminster Buena Brea Placentia Yorba Linda Cypress Stanton Seal
    Seaside Harbor Wilmington Pedro Lomita Redondo Hermosa Manhattan Palos Verdes Rolling Rancho Cucamonga Upland
    Loma Colton Highland Muscatine Moreno Eastvale Jurupa Lake Elsinore Wildomar Canyon Dana Capistrano Ladera
    Lagunas Lagos Forest Aliso Foothill Ranch Silverado Modjeska Trabuco Orange Villa Park Chino Pomona Claremont
    Montclair Azusa Glendora Covina Wilson Smith Thacher Shoemaker Tiebout Jung Silkworth Hazard Seiberling Ignatia
    Alexander Dowling Bob Lois Anne Ebby Clarence Snyder Marty Mann Jim Burwell Nell Wing Hank Parkhurst Fitz Mayo
    Carter Lincoln Washington Kennedy King Luther Christ Jesus God Dios Lord Senor Señor Higher Power Poder Superior
    Thank Thanks Gracias Welcome Bienvenidos Happy Feliz Merry Good Bueno Best Mejor New Nuevo Old Viejo First
    Primer Second Segundo Third Tercer Fourth Fifth Last Ultimo Next Proximo Próximo Final Draft Borrador Approved
    Aprobado Motion Moved Seconded Passed Failed Carried Unanimous Unanimously Vote Votes Voting Votos Total Totals
    Yes No Si Sí Abstain Abstentions Abstenciones Present Presente Absent Ausente Excused Guest Guests Invitados
    Visitors Visitantes Zoom Online Hybrid Hibrido Híbrido Virtual Person Persona Free Gratis Open Abierta Closed
    Cerrada Women Mujeres Men Hombres Spirituality Espiritualidad Unidos Union Unión Fe Esperanza Amor Love Peace
    Paz Gratitude Gratitud Dinner Cena Lunch Almuerzo Breakfast Desayuno Potluck Picnic Dance Baile Party Fiesta
    Roundup Round Up Rally Retreat Retiro Marathon Maratón Servathon Servatón Heritage Day Dia Día Night Noche
    Morning Mañana Afternoon Tarde Evening Week Semana Month Mes Year Año Quarter Trimestre Hour Hora Minute Minuto
    Box Caja Book Books Libros Pamphlet Pamphlets Folleto Folletos Workbook Manual Handbook Guide Guía Guia Form
    Forms Formulario Flyer Volante Website Sitio Email Correo Phone Telefono Teléfono Address Direccion Dirección
    Jr Sr II III IV PhD MD Dr Mr Mrs Ms Miss Sra Sr Srta Don Doña Rev Father Padre Pastor Sister Hermana Brother
    Hermano Mother Madre Nurse Doctor Judge Juez Officer Oficial Sheriff Deputy Captain Sergeant Chaplain
    Capellan Capellán Warden Director Directora Manager Gerente Staff Personal Employee Empleado Volunteer
    Voluntario Coordinator Assistant Asistente President Presidente Vice Chairman Moderator Moderador Facilitator
    Host Hosts Anfitrion Anfitrión Founder Founders Fundadores Cofounder Co Pioneer Pioneers Old Timer Timers""".split()
)
ROLE_NEAR = re.compile(
    r"\b(dcmc|dcm|gsr|gsrs|cmcd|mcd|mcds|rsg|rsgs|delegate|delegado|delegada|alternate|suplente|alt\.|chair|"
    r"chairperson|co-?chair|coordinador|coordinadora|coordinator|secretary|secretario|secretaria|treasurer|tesorero|"
    r"tesorera|registrar|registrador|registradora|archivist|archivista|trustee|custodio|custodia|liaison|enlace|"
    r"webmaster|editor|editora|representative|representante|panel\s+\d+|district\s+\d+|distrito\s+\d+|"
    r"submitted|respectfully|sincerely|atentamente|cordialmente|presented by|presentad[oa] por|prepared by|"
    r"preparad[oa] por|moved by|seconded|secundad[oa]|motion by|moci[oó]n de|signed|firmado|contact|contacto|"
    r"call|llame|llamar|elected|electo|electa|elegid[oa]|nominated|nominad[oa]|resigned|renunci[oó]|in loving memory|"
    r"memoriam|passed away|falleci[oó])\b",
    re.I,
)
ATTENDANCE = re.compile(
    r"\b(present|in attendance|attendance|attendees|attending|members present|officers present|absent|excused|"
    r"roll call|presentes|asistentes|asistencia|pase de lista|ausentes)\s*[:\-–]",
    re.I,
)
_SURNAME = r"((?:Mc|Mac|O'|De\s|Del\s|De\sla\s|Van\s|Von\s|St\.\s)?[A-Z][a-záéíóúñü][A-Za-záéíóúñü'\-]{1,20})"
NAME_RX = re.compile(r"\b(" + "|".join(sorted(FIRST_NAMES, key=len, reverse=True)) + r")\s+" + _SURNAME + r"(?![\w.])")
NAME_RX_UP = re.compile(
    r"\b(" + "|".join(sorted((n.upper() for n in FIRST_NAMES), key=len, reverse=True)) + r")\s+([A-Z][A-Z'\-]{2,20})(?![\w.])"
)
# More words that follow a first name in minutes and reports but are not surnames.
NOT_SURNAME |= set(
    """Concept Concepts Concepto Conceptos Introductions Introduction Introducciones Presentaciones Presentacion
    Presentación Tree Alt Call Dist Correctional United Birthdays Birthday Cumpleaños Approval Aprobacion Aprobación
    Review Reported Reports Reporto Reportó Informo Informó Attendance Visitor Visitors Visitante Visitantes GSR GSRs
    RSG RSGs DCM DCMC MCD CMCD Standing Statement Statements Twain Imperial Alcoholico Alcohólico Alcoholica
    Alcohólica Alcoholic Preamble Preámbulo Preambulo Methodist Motioned Moved Seconded AND THE Lane Banquet Sub
    Discussed Sur Alterno Alterna Eboard Declaración Declaracion Declaration Habra Tambien También Español Espanol
    Pre-Conference Opening Closing Prayer Oración Oracion Moment Silence Reading Lectura Lecturas Asked Announced
    Announcements Anuncios Agreed Explained Shared Compartio Compartió Thanked Agradecio Agradeció Gave Read Leyo
    Leyó Said Dijo Spoke Hablo Habló Will Would Was Is Has Had And Also Then Please Por Para Con Sin Que Quien
    Who What When Where Why How Old Business New Nuevo Asuntos Pendientes Treasurer's Chair's Delegate's
    Correction Corrections Committee's Questions Preguntas Answers Respuestas Comments Comentarios Update Updates
    Pending Tabled Withdrawn Amended Amendment Enmienda Floor Piso Quorum Quórum Roll Count Total Ayes Nays
    Opposed Abstained Favor Contra Recess Receso Break Adjourned Adjourn Clausura Cierre Lunch Agenda Orden
    Del Día Dia Minute Speaker Orador Oradora Moderador Moderator Panelist Panelists Panelistas Mentor Mentors
    Sponsor Sponsee Ahijado Padrino Madrina Liaisons Enlaces Coordinators Coordinadores Interpreter Interprete
    Intérprete Translation Traduccion Traducción Translator Traductor Archivist Archivista Registrar's Webmaster
    Webservant Editor Tech Technology Website Newsletter Boletin Boletín Grapevine Viña Literatura Pamphlets
    Joshua Elsinore Viejo Juan Capistrano Clemente Diego Imperial Bernardino Gabriel Jacinto Fernando Pedro
    Paula Rosa Monica Barbara Ana Cruz Marcos Luis Obispo Enlace Special Especial San Van Von Into Date
    Accounts Payable Receivable Cuentas Pagar Cobrar Remarks Highlights Notes Notas Summary Resumen Election Elecciones
    Elections Orientation Orientacion Orientación Toolkit Kit Packet Paquete Training Capacitacion Capacitación
    Hello Hi Hola Dear Querido Queridos Querida Estimado Estimados Estimada Saludos Greetings Permanente Permanentes
    Nuevos Nuevas Nuevo Nueva Docusign DocuSign Confidential Confidencial Vacancy Vacant Vacante Vacantes Quarterly
    Monthly Annual Anual Mensual Trimestral Filled Open Abierto Position Puesto Puestos Positions Rotation Rotación
    Signature Firma Envelope Completed Submitted Approved Pending Sharing Compartimiento Workshop Mock Simulacro
    Committee Comité Comite Report Informe Contributions Contribuciones Budget Treasury Tesorería Account Cuenta
    See Vivir Eterna TreasurerAR TreasurerAP Jan Feb Mar Apr Jun Jul Aug Sep Sept Oct Nov Dec Mon Tue Wed Thu
    Fri Sat Sun Ene Abr Ago Dic""".split()
)
NOT_SURNAME_LC = {w.lower() for w in NOT_SURNAME}
PLACE_BEFORE = re.compile(r"(?:\bSan|\bSanta|\bSan\s+Juan|\bLos|\bLas|\bLa|\bEl|\bSt\.?|\bSaint|\bMount|\bMt\.?|\bLake|\bFort|\bPort|\bPalm|\bRancho|\bLaguna)\s*$")
VENUE_AFTER = re.compile(
    r"^\s*(?:Methodist|Church|Iglesia|Center|Centre|Hall|Park|School|Hospital|College|University|Memorial|Lutheran|"
    r"Presbyterian|Catholic|Baptist|Episcopal|Elementary|High|Middle|Library|Room|Blvd|Boulevard|Street|St\.?|Avenue|Ave\.?|"
    r"Road|Rd\.?|Drive|Dr\.?|Way|Lane|Ln\.?|Hotel|Inn|Club|Fellowship|Group|Grupo|Foundation|Fund|Award|Prize|House|Home|"
    r"Institute|Instituto|Program|Programa|Clinic|Medical|Valley|Hills|Beach|Springs|Canyon|Tree|Unified|County|City)\b"
)
# People whose full names are public: non-alcoholic (Class A) trustees named as such in A.A. literature,
# non-alcoholic G.S.O. staff named in G.S.O. reports, A.A.'s co-founders and other historic figures.
# (A G.S.O. staff member listed here is not the same person as any trusted servant listed by initial.) Everyone else appears as first name + last initial.
PUBLIC_PEOPLE = {
    "bill wilson", "lois wilson", "bob smith", "anne smith", "ebby thacher", "sam shoemaker", "carl jung",
    "william silkworth", "harry tiebout", "jack alexander", "henrietta seiberling", "rowland hazard", "mark twain",
    "bernard smith", "jack norris", "john norris", "michael alexander", "jim estelle", "michele grinberg",
    "michelle mirza", "nancy mccarthy", "ivan lemelle", "rogelio flores", "linda chezem", "david morris",
    "leslie backus", "paul konigstein", "john fromson", "elaine mcdowell", "bruce hartley", "lola ibrahim",
    "beverly jones-king", "ann karam", "kevin prior", "leonard blumenthal", "terry bedient", "terrance bedient",
    "christine carpenter", "april hegner", "amy filiatreau", "janet bryan", "jonathan lobo", "marissa sblendorio",
    "alexandra rosenman", "john bealer", "kyle zaczek", "peter luongo", "allen ault", "al mooney", "martin luther",
}
CLASS_A_LABEL = re.compile(
    r"(?:Class\s*A|Clase\s*A|non-?alcoholic|no\s*alcoh[oó]lic[oa])", re.I
)
_LOWER_WORD = re.compile(r"(?<![A-Za-z])([a-záéíóúñü]{3,})(?![A-Za-z])")


def doc_vocabulary(text):
    """Lower-case words used in a document: a capitalised word that also appears in lower case is not a surname."""
    return set(_LOWER_WORD.findall(text))


ROLE_TITLE = (
    r"(?:Delegate|Delegado|Delegada|Alt\.?\s*Delegate|Alternate\s+Delegate|Delegad[oa]\s+Suplente|Suplente|"
    r"Chair(?:person|man|woman)?|Co-?Chair|Area\s+Chair|Secretary|Recording\s+Secretary|Secretari[oa]|"
    r"Treasurer(?:\s+(?:Accounts\s+(?:Payable|Receivable)|A/?P|A/?R))?|Tesorer[oa]|Registrar|Registrador[a]?|"
    r"Archivist|Archivista|DCMC|DCM|GSR|CMCD|MCD|RSG|Coordinator|Coordinador[a]?|Liaison|Enlace|Editor[a]?|"
    r"Webmaster|Trustee|Custodi[oa]|Treasurer\s+AP|Treasurer\s+AR|Committee\s+Chair|Coordinador[a]?\s+del?\s+Comit[eé])"
)
# "Secretary\nMaryka Lastname", "DCMC: Pat Lastname", "Lastname, Pat — Treasurer" — a capitalised pair right
# next to a service title, whatever the first name (rosters, signatures, contact tables).
ROLE_PAIR = re.compile(
    r"\b" + ROLE_TITLE + r"\b[ \t]*[:\-–,]?[ \t]*(?:\n[ \t]*)?"
    r"([A-Z][a-záéíóúñü]{2,20})[ \t]+((?:(?:de|del|de la|van|von|da|di|la|le)[ \t]+)?(?:Mc|Mac|O')?[A-Z][a-záéíóúñü][A-Za-záéíóúñü'\-]{1,20})(?![\w.])"
)
ROLE_PAIR_AFTER = re.compile(
    r"(?<![\w.])([A-Z][a-záéíóúñü]{2,20})[ \t]+((?:(?:de|del|de la|van|von|da|di|la|le)[ \t]+)?(?:Mc|Mac|O')?[A-Z][a-záéíóúñü][A-Za-záéíóúñü'\-]{1,20})"
    r"[ \t]*(?:,|–|-|\(|\n)[ \t]*" + ROLE_TITLE + r"\b"
)


def _skip_pair(text, m, first, last, vocab, strong):
    core = re.sub(r"^(?:(?:de|del|de la|van|von|da|di|la|le)\s+)?(Mc|Mac|O')?", "", last, flags=re.I)
    if not core or core.lower() in NOT_SURNAME_LC or last.lower() in NOT_SURNAME_LC or first.lower() in NOT_SURNAME_LC or "--" in last:
        return True
    if core.capitalize() in FIRST_NAMES and core.capitalize() not in ("Lee", "Marshall", "Grant", "James", "Thomas", "Allen", "Martin", "Nelson", "Russell", "Wayne", "High"):
        return True  # "Maria Elena", "Juan Carlos": two first names
    if len(core) < 3:
        return True
    if not strong and core.lower() in vocab:
        return True  # "Amy Concept…" — the word is used in lower case elsewhere: not a surname
    if f"{first} {last}".lower() in PUBLIC_PEOPLE:
        return True
    if PLACE_BEFORE.search(text[max(0, m.start(1) - 12) : m.start(1)]) or VENUE_AFTER.match(text[m.end(2) : m.end(2) + 14]):
        return True  # "Santa Ana Central Office", "Wesley United Methodist", "Jesse High School"
    around = text[max(0, m.start(1) - 60) : m.end(2) + 60]
    if CLASS_A_LABEL.search(around) and re.search(r"trustee|custodi|class|clase|alcoh", around, re.I):
        return True  # a non-alcoholic (Class A) trustee, named in full by A.A. itself
    return False


def name_hits(text, vocab=frozenset()):
    """(start, end, first, last) of 'First Lastname' next to a service title, near a role word, in an
    attendance list ("Present:") or above a signature."""
    blocks = [(m.end(), m.end() + 900) for m in ATTENDANCE.finditer(text)]
    out = []
    seen = set()
    # 1. Any capitalised pair right next to a service title (strong: the first name need not be a known one).
    for rx in (ROLE_PAIR, ROLE_PAIR_AFTER):
        for m in rx.finditer(text):
            first, last = m.group(1), m.group(2).strip()
            if first in FIRST_NAMES or first.upper() == first:
                continue  # known first names are handled below
            core = last.split()[-1].lower()
            if first.lower() in NOT_SURNAME_LC or first.lower() in vocab or core in vocab:
                continue  # "Treasurer Accounts Payable", "Chair Opening Remarks": ordinary words, not a name
            if _skip_pair(text, m, first, last, vocab, True):
                continue
            seen.add(m.start(1))
            out.append((m.start(1), m.end(2), first, last))
    # 2. A known first name + a surname near a role word or in an attendance list.
    for rx in (NAME_RX, NAME_RX_UP):
        for m in rx.finditer(text):
            if m.start(1) in seen:
                continue
            first, last = m.group(1), m.group(2).strip()
            in_block = any(a <= m.start() <= b for a, b in blocks)
            ctx = text[max(0, m.start() - 70) : m.end() + 70]
            strong = in_block or bool(re.search(ROLE_TITLE, text[max(0, m.start() - 30) : m.end() + 30]))
            if not (in_block or ROLE_NEAR.search(ctx)):
                continue
            if _skip_pair(text, m, first, last, vocab, strong):
                continue
            out.append((m.start(1), m.end(2), first, last))
    return out


# ------------------------------------------------------------------ text extraction
def _text_pymupdf(path):
    import fitz  # PyMuPDF

    fitz.TOOLS.mupdf_display_errors(False)
    pages = []
    with fitz.open(path) as d:
        if d.needs_pass:
            d.authenticate("")
        for p in d:
            try:
                t = p.get_text("text") or ""
            except Exception:
                t = ""
            try:
                imgs = len(p.get_images(full=False))
            except Exception:
                imgs = 0
            pages.append({"t": t, "imgs": imgs})
    return pages


def _text_pypdf(path, n):
    import pypdf

    out = [""] * n
    try:
        rd = pypdf.PdfReader(path, strict=False)
        if rd.is_encrypted:
            try:
                rd.decrypt("")
            except Exception:
                return out
        for i, p in enumerate(rd.pages):
            try:
                t = p.extract_text() or ""
            except Exception:
                t = ""
            if i < n:
                out[i] = t
            else:
                out.append(t)
    except Exception:
        pass
    return out


def _text_docx(path):
    import docx

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    for s in d.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            try:
                parts.extend(p.text for p in hf.paragraphs)
            except Exception:
                pass
    # Text boxes and anything else inside the XML
    try:
        xml = d.element.xml
        parts.append(" ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)))
    except Exception:
        pass
    return "\n".join(parts)


def _text_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            parts.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def _text_pptx(path):
    """Every text run of every slide and speaker note (read straight from the file's XML)."""
    import html
    import zipfile

    parts = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if re.match(r"ppt/(slides|notesSlides)/[^/]+\.xml$", n):
                xml = z.read(n).decode("utf-8", "ignore")
                parts.append(" ".join(html.unescape(t) for t in re.findall(r"<a:t>([^<]*)</a:t>", xml)))
    return "\n".join(parts)


def _text_binary(path):
    """Old binary Office files (.doc, .xls, .ppt): the readable text runs (8-bit and UTF-16)."""
    data = open(path, "rb").read()
    runs = [m.group(0).decode("cp1252", "ignore") for m in re.finditer(rb"[\x20-\x7e\xa0-\xff\r\n\t]{4,}", data)]
    runs += [m.group(0).decode("utf-16-le", "ignore") for m in re.finditer(rb"(?:[\x20-\x7e\xa0-\xff\r\n\t]\x00){4,}", data)]
    return "\n".join(runs)


def extract(path):
    """{ kind, pages: [ {a: pymupdf text, b: pypdf text, imgs} ], error }"""
    ext = path.rsplit(".", 1)[-1].lower()
    try:
        if ext == "pdf":
            pm = _text_pymupdf(path)
            pp = _text_pypdf(path, len(pm))
            pages = [{"a": pm[i]["t"], "b": pp[i] if i < len(pp) else "", "imgs": pm[i]["imgs"]} for i in range(len(pm))]
            return {"kind": "pdf", "pages": pages}
        if ext == "docx":
            return {"kind": "docx", "pages": [{"a": _text_docx(path), "b": "", "imgs": 0}]}
        if ext == "xlsx":
            return {"kind": "xlsx", "pages": [{"a": _text_xlsx(path), "b": "", "imgs": 0}]}
        if ext == "pptx":
            return {"kind": "pptx", "pages": [{"a": _text_pptx(path), "b": "", "imgs": 0}]}
        if ext in ("doc", "xls", "ppt", "rtf", "txt"):
            return {"kind": ext, "pages": [{"a": _text_binary(path), "b": "", "imgs": 0}]}
        return {"kind": ext, "pages": [], "error": "unsupported file type (open it and check it by hand)"}
    except ModuleNotFoundError as e:
        return {"kind": ext, "pages": [], "error": f"missing Python package: {e.name}"}
    except Exception as e:  # damaged or password-protected file
        return {"kind": ext, "pages": [], "error": type(e).__name__}


def _norm_space(t):
    t = unicodedata.normalize("NFKC", t or "")
    t = t.replace("\u00a0", " ").replace("\u200b", "")
    return re.sub(r"[ \t]+", " ", t)


# ------------------------------------------------------------------ the scan itself
COMMON_FILE = os.path.join(CACHE_DIR, "scan-documents-common.json")
COMMON_MIN_DOCS = 4  # an address printed in this many documents is a meeting place or an office


class Rules:
    def __init__(self):
        self.allow_emails, self.allow_phones = load_allow()
        self.venues = load_venues()
        try:
            self.common = frozenset(json.load(open(COMMON_FILE, encoding="utf-8")))
        except Exception:
            self.common = frozenset()


def find_all(text, rules, vocab=frozenset(), common=frozenset()):
    """Findings in one text: a dict of kind → list of exact strings (kept in memory only)."""
    t = _norm_space(text)
    f = {"email": [], "phone": [], "sobriety": [], "address": [], "name": [], "_addr": []}
    for m in EMAIL.finditer(join_split_emails(t)):
        if email_is_personal(m.group(0), rules.allow_emails):
            f["email"].append(m.group(0))
    for d, raw in phone_hits(t, rules.allow_phones):
        f["phone"].append(raw)
    for a, b in sobriety_hits(t):
        f["sobriety"].append(t[a:b])
    for a, b, key, personal in address_hits(t, rules.venues, common):
        f["address"].append(t[a:b])
        f["_addr"].append((addr_hash(key), personal))
    for a, b, first, last in name_hits(t, vocab):
        f["name"].append(t[a:b])
    return f


def scan_extracted(ex, rules, ocr_text=None, use_common=False):
    """Counts per kind + per-page findings (strings stay in memory)."""
    per_page = []
    unverified = 0
    chars = 0
    alltext = " ".join(pg["a"] + " " + pg["b"] for pg in ex["pages"]) + " " + " ".join((ocr_text or {}).values())
    alltext = re.sub(r"\S+@\S+|https?://\S+|www\.\S+", " ", alltext)  # names inside e-mail addresses are not "words"
    vocab = doc_vocabulary(alltext)
    for i, pg in enumerate(ex["pages"]):
        a, b = pg["a"], pg["b"]
        extra = (ocr_text or {}).get(i, "")
        n_chars = max(len(re.sub(r"\s", "", a)), len(re.sub(r"\s", "", b)))
        chars += n_chars
        if ex["kind"] == "pdf" and n_chars < 25 and pg["imgs"] and len(re.sub(r"\s", "", extra)) < 25:
            unverified += 1
        common = rules.common if use_common else frozenset()
        fa = find_all(a, rules, vocab, common)
        fb = find_all(b, rules, vocab, common) if b else {k: [] for k in fa}
        fo = find_all(extra, rules, vocab, common) if extra else {k: [] for k in fa}
        merged = {k: sorted(set(fa[k]) | set(fb[k]) | set(fo[k])) for k in fa}
        merged["_only_pypdf"] = sorted({x for k in fa if k != "_addr" for x in fb[k] if x not in fa[k] and _squash(x) not in _squash(a)})
        merged["_ocr"] = sorted({x for k in fa if k != "_addr" for x in fo[k]})
        per_page.append(merged)
    counts = {}
    for k in ("email", "phone", "sobriety", "address", "name"):
        counts[k] = len({_squash(x) for p in per_page for x in p[k]})
    if use_common:
        counts["address"] = len({_squash(x) for p in per_page for x in p["address"]})
    counts["unverified_pages"] = unverified
    counts["pages"] = len(ex["pages"])
    counts["chars"] = chars
    addr = {}
    for p in per_page:
        for h, personal in p["_addr"]:
            addr[h] = addr.get(h, False) or personal
    counts["addr"] = sorted([h, int(v)] for h, v in addr.items())  # one-way hashes only
    return counts, per_page


def _squash(s):
    return re.sub(r"[\s\-–.()/]", "", (s or "").lower())


def flags_of(counts, error=None):
    f = []
    if error:
        f.append("unverified")
    if counts.get("email") or counts.get("phone") or counts.get("address"):
        f.append("personal-contact")
    if counts.get("name"):
        f.append("full-name")
    if counts.get("sobriety"):
        f.append("sobriety-date")
    if counts.get("unverified_pages"):
        f.append("unverified")
    return sorted(set(f))


# ------------------------------------------------------------------ OCR (Windows 10/11, optional)
OCR_PS = r"""
param([string]$ListFile)
Add-Type -AssemblyName System.Runtime.WindowsRuntime
$null = [Windows.Storage.StorageFile, Windows.Storage, ContentType = WindowsRuntime]
$null = [Windows.Media.Ocr.OcrEngine, Windows.Foundation, ContentType = WindowsRuntime]
$null = [Windows.Graphics.Imaging.BitmapDecoder, Windows.Foundation, ContentType = WindowsRuntime]
$asTaskGeneric = ([System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq 'AsTask' -and $_.GetParameters().Count -eq 1 -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1' })[0]
function Await($WinRtTask, $ResultType) { $t = $asTaskGeneric.MakeGenericMethod($ResultType).Invoke($null, @($WinRtTask)); $t.Wait(-1) | Out-Null; $t.Result }
$engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
foreach ($p in Get-Content -Encoding UTF8 $ListFile) {
  try {
    $file = Await ([Windows.Storage.StorageFile]::GetFileFromPathAsync($p)) ([Windows.Storage.StorageFile])
    $stream = Await ($file.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
    $decoder = Await ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream)) ([Windows.Graphics.Imaging.BitmapDecoder])
    $bitmap = Await ($decoder.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
    $result = Await ($engine.RecognizeAsync($bitmap)) ([Windows.Media.Ocr.OcrResult])
    [System.IO.File]::WriteAllText($p + ".txt", $result.Text, [System.Text.Encoding]::UTF8)
    $stream.Dispose()
  } catch { }
}
"""


def ocr_pages(path, page_nums, workdir):
    """{page index: text} for the given pages, via the Windows OCR engine. {} when OCR is unavailable."""
    if os.name != "nt" or not page_nums:
        return {}
    import fitz

    pngs = {}
    with fitz.open(path) as d:
        for i in page_nums:
            p = d[i]
            z = min(2400 / max(p.rect.width, 1), 4)
            out = os.path.join(workdir, f"{hashlib.sha1((path + str(i)).encode()).hexdigest()[:16]}.png")
            p.get_pixmap(matrix=fitz.Matrix(z, z)).save(out)
            pngs[i] = out
    lst = os.path.join(workdir, "ocr-list.txt")
    ps1 = os.path.join(workdir, "ocr.ps1")
    open(lst, "w", encoding="utf-8").write("\n".join(pngs.values()))
    open(ps1, "w", encoding="utf-8").write(OCR_PS)
    try:
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", ps1, lst], capture_output=True, timeout=600)
    except Exception:
        return {}
    out = {}
    for i, png in pngs.items():
        t = png + ".txt"
        if os.path.exists(t):
            out[i] = open(t, encoding="utf-8-sig").read()
            os.remove(t)
        os.remove(png)
    return out


# ------------------------------------------------------------------ per-file worker
_RULES = None


def get_rules():
    global _RULES
    if _RULES is None:
        _RULES = Rules()
    return _RULES


def scan_path(path, use_ocr=False, use_common=False):
    rules = get_rules()
    ex = extract(path)
    ocr = {}
    if use_ocr and ex["kind"] == "pdf":
        need = [i for i, pg in enumerate(ex["pages"]) if pg["imgs"] and max(len(re.sub(r"\s", "", pg["a"])), len(re.sub(r"\s", "", pg["b"]))) < 25]
        if need:
            with tempfile.TemporaryDirectory() as td:
                ocr = ocr_pages(path, need, td)
    counts, _ = scan_extracted(ex, rules, ocr, use_common)
    if use_common:
        counts["address"] = len({h for h, personal in counts.get("addr", []) if personal or h not in rules.common})
    if ocr:
        counts["ocr_pages"] = len(ocr)
    return {"counts": counts, "flags": flags_of(counts, ex.get("error")), "error": ex.get("error", ""), "kind": ex["kind"]}


def _worker(args):
    path, use_ocr = args
    try:
        return path, scan_path(path, use_ocr)
    except Exception as e:
        return path, {"counts": {}, "flags": ["unverified"], "error": type(e).__name__, "kind": ""}


# ------------------------------------------------------------------ redaction (PDF)
def _rects_for(page, needle):
    import fitz

    rects = list(page.search_for(needle))
    if rects:
        return rects
    # Word-by-word match (the text may be split over lines or have odd spacing).
    toks = [w for w in re.split(r"\s+", needle.strip()) if w]
    if not toks:
        return []
    words = page.get_text("words")
    clean = lambda s: re.sub(r"[^\w@.+\-]", "", s).lower()
    want = [clean(t) for t in toks]
    out = []
    for i in range(len(words)):
        if clean(words[i][4]) != want[0] and not (len(want) == 1 and want[0] in clean(words[i][4])):
            continue
        ok = True
        for j in range(1, len(want)):
            if i + j >= len(words) or clean(words[i + j][4]) != want[j]:
                ok = False
                break
        if ok:
            for j in range(len(want)):
                out.append(fitz.Rect(words[i + j][:4]))
    if out:
        return out
    # The text is broken over two or more lines ("name@yahoo.c" + "om"): join neighbouring words.
    target = "".join(want)
    for i in range(len(words)):
        acc = ""
        for j in range(i, min(i + 5, len(words))):
            acc += clean(words[j][4])
            if not target.startswith(acc[: len(target)]) and not acc.startswith(target):
                break
            if len(acc) >= len(target):
                if acc.startswith(target):
                    out.extend(fitz.Rect(words[k][:4]) for k in range(i, j + 1))
                break
    return out


def redact_pdf(src, dst, rules=None):
    """Write a cleaned copy of src to dst. Returns (hits, not_located) — counts only."""
    import fitz

    rules = rules or get_rules()
    ex = extract(src)
    _, per_page = scan_extracted(ex, rules, use_common=True)
    hits = 0
    missing = 0
    with fitz.open(src) as d:
        for i, page in enumerate(d):
            f = per_page[i] if i < len(per_page) else {}
            todo = []
            for k in ("email", "phone", "sobriety", "address"):
                for s in f.get(k, []):
                    todo.append((s, "[removed]"))
            for s in f.get("name", []):
                parts = s.split()
                todo.append((s, f"{parts[0].capitalize() if parts[0].isupper() else parts[0]} {parts[-1][0]}."))
            for s, repl in sorted(todo, key=lambda x: -len(x[0])):
                rects = _rects_for(page, s)
                if not rects and s.upper() != s:
                    rects = _rects_for(page, s.upper())
                if not rects:
                    missing += 1
                    continue
                for r in rects:
                    page.add_redact_annot(r, text=repl, fontsize=0, fill=(1, 1, 1))
                    hits += 1
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)
        d.set_metadata({"title": "", "author": "", "subject": "", "keywords": "", "creator": "", "producer": ""})
        d.del_xml_metadata()  # the XMP packet keeps its own copy of the author's name (dc:creator)
        d.save(dst, garbage=4, deflate=True)
    return hits, missing


# ------------------------------------------------------------------ CLI
def file_times(p):
    st = os.stat(p)
    return max(st.st_mtime, getattr(st, "st_birthtime", st.st_ctime))


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):  # Windows consoles (cp1252) cannot print every character
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--ids", help="comma-separated row ids, or @file with one id per line")
    ap.add_argument("--file", help="scan one local file (not from the index)")
    ap.add_argument("--redacted", action="store_true", help="only rows whose id ends with -redacted")
    ap.add_argument("--since", help="only files created or changed since this local time (YYYY-MM-DDTHH:MM)")
    ap.add_argument("--all-rows", action="store_true", help="also rows with publish = review / no")
    ap.add_argument("--ocr", action="store_true", help="OCR pages without a text layer (Windows 10/11)")
    ap.add_argument("--json", default=os.path.join(CACHE_DIR, "scan-documents.json"), help="where to write the per-document counts")
    ap.add_argument("--workers", type=int, default=max(2, (os.cpu_count() or 4) - 2))
    ap.add_argument("--no-cache", action="store_true")
    ap.add_argument("--hold", action="store_true", help="move flagged publish=yes rows to data/documents/held.csv")
    ap.add_argument("--redact", nargs=2, metavar=("IN", "OUT"), help="write a cleaned copy of a PDF, then re-scan it")
    ap.add_argument("--drive", help="your local copy of the Area Drive folder (default: the MSCA09_DRIVE environment variable)")
    a = ap.parse_args(argv)
    global DRIVE
    if a.drive:
        DRIVE = a.drive.rstrip("/\\")

    if a.redact:
        rules = Rules()
        hits, missing = redact_pdf(a.redact[0], a.redact[1], rules)
        res = scan_path(a.redact[1], a.ocr, use_common=True)
        print(f"redacted {hits} place(s); {missing} finding(s) could not be located on the page")
        res["counts"].pop("addr", None)
        print("after redaction:", json.dumps(res["counts"]), "flags:", ",".join(res["flags"]) or "none")
        return 0 if not res["flags"] else 1

    if a.file:
        res = scan_path(a.file, a.ocr, use_common=True)
        res["counts"].pop("addr", None)
        print(json.dumps({"file": os.path.basename(a.file), **res}, indent=1))
        return 1 if res["flags"] else 0

    if not DRIVE or not os.path.isdir(DRIVE):
        print("Set MSCA09_DRIVE (or --drive) to your Google Drive for Desktop copy of the Area folder, the one that\n"
              "contains docs/, e.g.  MSCA09_DRIVE=\"G:/My Drive/MSCA09AA\"" + (f" (not found: {DRIVE})" if DRIVE else ""), file=sys.stderr)
        return 2
    rows = []
    for f in index_files():
        rows.extend(read_rows(f)[1])
    if a.ids and a.ids.startswith("@"):  # --ids @list.txt : ids from a file (comma or line separated)
        a.ids = re.sub(r"\s+", ",", open(a.ids[1:], encoding="utf-8").read().strip())
    ids = set(filter(None, a.ids.split(","))) if a.ids else None
    since = dt.datetime.fromisoformat(a.since).timestamp() if a.since else None
    todo = []
    missing = []
    for r in rows:
        if ids and r["id"] not in ids:
            continue
        if not a.all_rows and not ids and r.get("publish", "").strip().lower() != "yes":
            continue
        if a.redacted and not r["id"].endswith("-redacted"):
            continue
        dp = r.get("drive_path", "").strip()
        if not dp:
            continue
        p = os.path.join(DRIVE, dp)
        if not os.path.exists(p):
            missing.append(r["id"])
            continue
        if since and file_times(p) < since:
            continue
        todo.append((r, p))

    cache_file = os.path.join(CACHE_DIR, "scan-documents-cache.json")
    cache = {}
    if not a.no_cache and os.path.exists(cache_file):
        try:
            cache = json.load(open(cache_file, encoding="utf-8"))
        except Exception:
            cache = {}
    results = {}
    jobs = []
    for r, p in todo:
        st = os.stat(p)
        key = f"{SCAN_VERSION}|{int(a.ocr)}|{p}|{st.st_size}|{int(st.st_mtime)}"
        if key in cache:
            results[r["id"]] = cache[key]
        else:
            jobs.append((r["id"], p, key))
    if jobs:
        print(f"scanning {len(jobs)} file(s) ({len(todo) - len(jobs)} unchanged, from cache)...", file=sys.stderr)
        by_path = {}
        for rid, p, key in jobs:
            by_path.setdefault(p, []).append((rid, key))
        done = 0
        with cf.ProcessPoolExecutor(max_workers=a.workers) as ex:
            for p, res in ex.map(_worker, [(p, a.ocr) for p in by_path], chunksize=1):
                for rid, key in by_path[p]:
                    results[rid] = res
                    cache[key] = res
                done += 1
                if done % 100 == 0:
                    print(f"  {done}/{len(by_path)}", file=sys.stderr)
                    os.makedirs(CACHE_DIR, exist_ok=True)
                    json.dump(cache, open(cache_file, "w", encoding="utf-8"))
        os.makedirs(CACHE_DIR, exist_ok=True)
        json.dump(cache, open(cache_file, "w", encoding="utf-8"))

    # Addresses printed in many documents are meeting places or offices: count each document's addresses
    # without them (unless the text says it is someone's home). Hashes only — nothing readable is stored.
    freq = {}
    for r, p in todo:
        for h, personal in (results.get(r["id"]) or {}).get("counts", {}).get("addr", []):
            if not personal:
                freq[h] = freq.get(h, 0) + 1
    # Remembered between runs (.cache/scan-documents-common.json): holding documents back must not turn a
    # meeting place into a "rare" address in the next run. Delete that file to start afresh.
    common = {h for h, n in freq.items() if n >= COMMON_MIN_DOCS} | set(get_rules().common)
    if not (a.ids or a.since or a.redacted or a.all_rows):  # a full run: remember them for later runs
        os.makedirs(CACHE_DIR, exist_ok=True)
        json.dump(sorted(common), open(COMMON_FILE, "w", encoding="utf-8"))
    for r, p in todo:
        res = results.get(r["id"])
        if not res or not res.get("counts"):
            continue
        c = dict(res["counts"])
        c["address"] = len({h for h, personal in c.get("addr", []) if personal or h not in common})
        results[r["id"]] = {**res, "counts": c, "flags": flags_of(c, res.get("error"))}

    out = []
    for r, p in todo:
        res = results.get(r["id"])
        if not res:
            continue
        counts = {k: v for k, v in res["counts"].items() if k != "addr"}
        out.append({"id": r["id"], "publish": r.get("publish", ""), "file": os.path.relpath(r["_file"], REPO).replace("\\", "/"),
                    "flags": res["flags"], "counts": counts, "error": res.get("error", "")})
    os.makedirs(os.path.dirname(os.path.abspath(a.json)), exist_ok=True)
    json.dump({"generated": dt.datetime.now().isoformat(timespec="seconds"), "drive": DRIVE, "scanned": len(out),
               "missing_files": missing, "documents": out}, open(a.json, "w", encoding="utf-8"), indent=1)

    flagged = [o for o in out if o["flags"]]
    by_flag = {}
    for o in flagged:
        for fl in o["flags"]:
            by_flag[fl] = by_flag.get(fl, 0) + 1
    print(f"scanned {len(out)} document(s); {len(flagged)} with findings; {len(missing)} row(s) whose file is not in your Drive copy")
    for fl, n in sorted(by_flag.items()):
        print(f"  {fl:18s} {n}")
    for o in flagged[:400]:
        c = o["counts"]
        bits = ", ".join(f"{k} {c[k]}" for k in ("email", "phone", "address", "sobriety", "name", "unverified_pages") if c.get(k))
        print(f"  - {o['id']}  [{o['publish']}]  {bits}{'  (' + o['error'] + ')' if o['error'] else ''}")
    print(f"details (counts only): {os.path.relpath(a.json, REPO) if a.json.startswith(REPO) else a.json}")

    if a.hold and flagged:
        n = hold_rows({o["id"]: o["flags"] for o in flagged if o["publish"].lower() == "yes"})
        print(f"moved {n} row(s) to data/documents/held.csv (publish = review)")
    return 1 if any(o["publish"].lower() == "yes" for o in flagged) else 0


def hold_rows(id_flags):
    """Move rows to data/documents/held.csv with publish = review and the anonymity flag."""
    d = os.path.join(DATA, "documents")
    held = os.path.join(d, "held.csv")
    if not os.path.isdir(d) or not os.path.exists(held):
        print("data/documents/held.csv not found — hold the rows by hand (publish = review).", file=sys.stderr)
        return 0
    moved = []
    for f in index_files():
        if os.path.samefile(f, held):
            continue
        with open(f, encoding="utf-8-sig", newline="") as fh:
            lines = list(csv.reader(fh))
        header = lines[0]
        keep = [header]
        for cells in lines[1:]:
            row = dict(zip(header, cells))
            if cells and not cells[0].startswith("#") and row.get("id") in id_flags:
                row["publish"] = "review"
                row["anonymity_flag"] = "; ".join(id_flags[row["id"]])
                row["notes"] = (row.get("notes", "") + " Held by scripts/scan-documents.py " + dt.date.today().isoformat()).strip()
                moved.append(row)
            else:
                keep.append(cells)
        if len(keep) != len(lines):
            with open(f, "w", encoding="utf-8", newline="") as fh:
                csv.writer(fh, lineterminator="\n").writerows(keep)
    if moved:
        with open(held, encoding="utf-8-sig", newline="") as fh:
            hheader = next(csv.reader(fh))
        with open(held, "a", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, lineterminator="\n")
            for row in moved:
                w.writerow([row.get(c, "") for c in hheader])
    return len(moved)


if __name__ == "__main__":
    sys.exit(main())
